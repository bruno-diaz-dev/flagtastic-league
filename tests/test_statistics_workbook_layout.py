"""Workbook layout regressions independent of database state."""
from io import BytesIO

import pytest
from openpyxl import Workbook

from services.statistics_import import (
    StatisticsFileError, parse_games_workbook, parse_statistics_workbook,
)


def workbook_bytes(repeated_header=False, empty_week=False, invalid=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Wk 1'
    header = ['Equipo', 'Categoria', 'Estadistica']
    sheet.append(header)
    labels = ['Intentos Pase', 'Completos Pase', 'Para % pases', 'Puntos Pase',
              '6 Puntos', '2 Puntos', '1 Puntos', 'Intercepciones Pase', 'TD',
              'Conv 1', 'Conv 2', 'Sacks', 'Tacleo', 'Intercepciones Def', 'Asistencia']
    for index, team in enumerate(('Home', 'Away')):
        if index and repeated_header:
            sheet.append(header)
        for label in labels:
            sheet.append([team, 'unknown' if invalid else 'U8', label,
                          7 if label in ('TD', 'Asistencia') else None])
    if empty_week:
        empty = workbook.create_sheet('Wk 2')
        empty.append(header)
        for label in labels:
            empty.append([None, None, label])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_repeated_headers_and_unused_week_templates_preserve_import():
    baseline = workbook_bytes()
    repeated = workbook_bytes(repeated_header=True, empty_week=True)
    assert parse_statistics_workbook(repeated) == parse_statistics_workbook(baseline)
    assert parse_games_workbook(repeated) == parse_games_workbook(baseline)
    assert set(parse_statistics_workbook(repeated)) == {1}


def test_invalid_populated_week_is_not_silently_skipped():
    with pytest.raises(StatisticsFileError, match='categoria desconocida'):
        parse_statistics_workbook(workbook_bytes(invalid=True, empty_week=True))


def test_selected_empty_week_still_reports_an_error():
    with pytest.raises(StatisticsFileError, match='jornada 2'):
        parse_statistics_workbook(workbook_bytes(empty_week=True), selected_week=2)


@pytest.mark.parametrize('label,expected', [
    ('U14 Varonil', ('varonil', 'u14')),
    ('U14 Femenil', ('femenil', 'u14')),
    ('Libre Mixto', ('mixto', 'libre')),
    ('Var Libre', ('varonil', 'libre')),
    ('FemU12', ('femenil', 'u12')),
])
def test_official_category_accepts_full_names_and_suffixes(label, expected):
    from services.statistics_import import _division_from_official_category
    assert _division_from_official_category(label, 'Wk 1', 182) == expected


def test_passing_points_are_not_personal_scoring_but_rushing_td_is():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Wk 1'
    labels = ['Intentos Pase', 'Completos Pase', 'Para % pases', 'Puntos Pase',
              '6 Puntos', '2 Puntos', '1 Puntos', 'Intercepciones Pase', 'TD',
              'Conv 1', 'Conv 2', 'Sacks', 'Tacleo', 'Intercepciones Def', 'Asistencia']
    values = {'6 Puntos': [5] * 7, '2 Puntos': [5] * 7, 'TD': [17, 5], 'Conv 2': [17], 'Asistencia': [5, 17]}
    sheet.append(['Equipo', 'Categoria', 'Estadistica'])
    for label in labels:
        sheet.append(['Diablos', 'Libre Femenil', label, *values.get(label, [])])
    for label in labels:
        sheet.append(['Rivals', 'Libre Femenil', label])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    rows = parse_statistics_workbook(output.getvalue())[1]
    by_number = {row['jersey_number']: row for row in rows}
    assert by_number[5]['points'] == 6  # Only her actual rushing touchdown.
    assert by_number[17]['points'] == 8


@pytest.mark.parametrize('attempts,completions', [(26, 14), (14, 14), (26, 0), (0, 0)])
def test_official_attempts_already_include_completed_passes(attempts, completions):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Wk 1'
    labels = ['Intentos Pase', 'Completos Pase', 'Para % pases', 'Puntos Pase',
              '6 Puntos', '2 Puntos', '1 Puntos', 'Intercepciones Pase', 'TD',
              'Conv 1', 'Conv 2', 'Sacks', 'Tacleo', 'Intercepciones Def', 'Asistencia']
    sheet.append(['Equipo', 'Categoria', 'Estadistica'])
    for label in labels:
        values = [8] * attempts if label == 'Intentos Pase' else [8] * completions if label == 'Para % pases' else [8] if label == 'Asistencia' else []
        sheet.append(['Storms', 'Mix Libre', label, *values])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    row = parse_statistics_workbook(output.getvalue())[1][0]
    assert row['passes_attempted'] == attempts
    assert row['passes_completed'] == completions


def test_official_completions_exceeding_attempts_are_rejected():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Wk 1'
    labels = ['Intentos Pase', 'Completos Pase', 'Para % pases', 'Puntos Pase',
              '6 Puntos', '2 Puntos', '1 Puntos', 'Intercepciones Pase', 'TD',
              'Conv 1', 'Conv 2', 'Sacks', 'Tacleo', 'Intercepciones Def', 'Asistencia']
    sheet.append(['Equipo', 'Categoria', 'Estadistica'])
    for label in labels:
        sheet.append(['Example', 'Mix Libre', label, *([8] if label == 'Para % pases' else [])])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(StatisticsFileError, match='completos exceden'):
        parse_statistics_workbook(output.getvalue())


@pytest.mark.parametrize('jersey', [0, 8])
def test_incomplete_official_passing_warns_and_preserves_other_metrics(jersey):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Wk 2'
    labels = ['Intentos Pase', 'Completos Pase', 'Para % pases', 'Puntos Pase',
              '6 Puntos', '2 Puntos', '1 Puntos', 'Intercepciones Pase', 'TD',
              'Conv 1', 'Conv 2', 'Sacks', 'Tacleo', 'Intercepciones Def', 'Asistencia']
    sheet.append(['Equipo', 'Categoria', 'Estadistica'])
    events = {'Completos Pase': [20], 'Para % pases': [jersey],
              'TD': [20], 'Tacleo': [jersey]}
    for label in labels:
        sheet.append(['Example', 'Fem 14', label, *events.get(label, [])])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    warnings = []
    rows = parse_statistics_workbook(output.getvalue(), warnings=warnings)[2]
    by_number = {row['jersey_number']: row for row in rows}
    assert by_number[20]['receptions'] == 1
    assert by_number[20]['points'] == 6
    assert by_number[jersey]['tackles'] == 1
    assert by_number[jersey]['passes_completed'] == 0
    assert by_number[jersey]['passes_attempted'] == 0
    assert by_number[jersey]['passing_incomplete'] is True
    assert len(warnings) == 1
    assert warnings[0]['source_completed'] == 1
    assert warnings[0]['source_attempted'] == 0
    assert warnings[0]['jersey_number'] == jersey
    assert warnings[0]['sheet'] == 'Wk 2'
    assert warnings[0]['row'] == 2
