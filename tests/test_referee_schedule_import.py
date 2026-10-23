"""Structured referee schedule parsing tests."""

from io import BytesIO

import pytest

from openpyxl import Workbook

from services.referee_schedule_import import RefereeScheduleImportError, parse_referee_schedule_file


GAMES = [{
    "id": 81,
    "week": 5,
    "field_number": 1,
    "start_time": "12:00",
    "home_team": {
        "id": 1, "name": "Nomadas U8", "branch": "mixto", "category": "u8"
    },
    "away_team": {
        "id": 2, "name": "Lobos U8", "branch": "mixto", "category": "u8"
    },
}]
REFEREES = [
    {"id": 11, "name": "Cameron Referee", "display_name": "Cam"},
    {"id": 12, "name": "Alan Judge", "display_name": "Alan"},
    {"id": 13, "name": "Miguel Field", "display_name": "Miguel"},
    {"id": 14, "name": "Allison Stats", "display_name": "Allison"},
]


def test_selected_week_ignores_history_and_empty_future_sheets():
    workbook = Workbook()
    workbook.remove(workbook.active)
    for week in range(1, 7):
        sheet = workbook.create_sheet(f"Semana {week}")
        sheet.append([f"Semana {week}"])
        sheet.append(["Hora", "Campo 1", None])
        if week <= 4:
            sheet.append(["12:00", "Nómadas U8", "Lobos U8"])
            sheet.append([None, "Cam / Alan / Miguel / Allison", None])
    content = BytesIO()
    workbook.save(content)
    games = [{**GAMES[0], "id": 80 + week, "week": week} for week in range(1, 5)]

    filtered = parse_referee_schedule_file(
        "historial.xlsx", content.getvalue(), games, REFEREES, selected_week=4
    )
    assert [proposal["game_id"] for proposal in filtered["proposals"]] == [84]
    assert filtered["selected_week"] == 4
    assert filtered["warnings"] == []
    assert filtered["unmatched"] == 0
    assert "Semana 4!" in filtered["raw_text"]
    assert "Semana 3!" not in filtered["raw_text"]
    # Missing historical matches must not generate warnings in a J4-only import.
    assert parse_referee_schedule_file(
        "historial.xlsx", content.getvalue(), [games[-1]], REFEREES, selected_week=4
    )["warnings"] == []
    assert parse_referee_schedule_file(
        "historial.xlsx", content.getvalue(), games, REFEREES
    )["matched"] == 4
    with pytest.raises(RefereeScheduleImportError, match="jornada 5"):
        parse_referee_schedule_file(
            "historial.xlsx", content.getvalue(), games, REFEREES, selected_week=5
        )


def test_csv_week_filter_excludes_other_rows_before_slot_validation():
    content = ("jornada,campo,hora,local,visitante,referee,down judge\n"
               "3,invalid,invalid,Otro,Lobos U8,No registrado,Otro\n"
               "4,1,12:00,Nómadas U8,Lobos U8,Cam,Alan\n").encode()
    result = parse_referee_schedule_file(
        "rol.csv", content, [{**GAMES[0], "week": 4}], REFEREES, selected_week=4
    )
    assert result["matched"] == 1
    assert result["warnings"] == []


def test_invited_friendly_teams_use_unique_exact_names_not_a_guessed_category():
    game = {**GAMES[0], "week": 4, "is_friendly": True,
            "home_team": {"id": 191, "name": "Sel AGS Var U17", "branch": "varonil", "category": "u18", "is_guest": True},
            "away_team": {"id": 192, "name": "Sel AGS Var U15", "branch": "varonil", "category": "u18", "is_guest": True}}
    content = ("jornada,campo,hora,local,visitante,referee,down judge\n"
               "4,1,12:00,Sel AGS Var U17,Sel AGS Var U15,Cam,Alan\n").encode()
    result = parse_referee_schedule_file("rol.csv", content, [game], REFEREES, selected_week=4)
    assert result["matched"] == 1
    assert result["warnings"] == []
    # U16 must not silently become the registered U17 guest, even at this slot.
    assert parse_referee_schedule_file(
        "rol.csv", content.replace(b"U17", b"U16"), [game], REFEREES, selected_week=4
    )["matched"] == 0


def test_analysis_endpoint_applies_week_filter_without_saving(monkeypatch):
    from fastapi.testclient import TestClient
    from main import app

    monkeypatch.setattr("routes.games.get_games", lambda: [
        {**GAMES[0], "id": week, "week": week} for week in [3, 4]
    ])
    monkeypatch.setattr("routes.games.get_referee_match_candidates", lambda: REFEREES)
    def no_writes(*args, **kwargs):
        pytest.fail("Analysis must never write assignments")
    monkeypatch.setattr("routes.games.apply_referee_schedule", no_writes)
    content = ("jornada,campo,hora,local,visitante,referee,down judge\n"
               "3,1,12:00,Nómadas U8,Lobos U8,Cam,Alan\n"
               "4,1,12:00,Nómadas U8,Lobos U8,Cam,Alan\n").encode()
    with TestClient(app) as client:
        response = client.post("/api/games/referee-schedule/analyze", data={"week": "4"},
                               files={"file": ("rol.csv", content, "text/csv")})
        assert response.status_code == 200
        assert response.json()["selected_week"] == 4
        assert [proposal["game_id"] for proposal in response.json()["proposals"]] == [4]
        for invalid in ["0", "-1", "4.5", "incorrecta"]:
            assert client.post("/api/games/referee-schedule/analyze", data={"week": invalid},
                               files={"file": ("rol.csv", content, "text/csv")}).status_code == 422


def test_xlsx_grid_reads_game_row_and_ordered_official_row():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Semana 5"
    sheet.append(["Semana 5"])
    sheet.append(["Hora", "Campo 1", None])
    sheet.append(["12:00", "Nómadas U8", "Lobos U8"])
    sheet.append([None, "Cam / Alan / Miguel / Allison", None])
    content = BytesIO()
    workbook.save(content)

    result = parse_referee_schedule_file(
        "rol-arbitros.xlsx", content.getvalue(), GAMES, REFEREES
    )

    assert result["matched"] == 1
    assert result["unmatched"] == 0
    assert result["warnings"] == []
    assert result["proposals"][0]["game_id"] == 81
    assert result["proposals"][0]["scheduled_time"] == "12:00"
    assert result["proposals"][0]["officials"] == [
        {"user_id": 11, "name": "Cam", "position": "referee"},
        {"user_id": 12, "name": "Alan", "position": "down_judge"},
        {"user_id": 13, "name": "Miguel", "position": "field_judge"},
        {"user_id": 14, "name": "Allison", "position": "statistician"},
    ]


def test_csv_table_accepts_explicit_role_columns_and_reports_unknown_names():
    content = (
        "jornada,campo,hora,local,visitante,referee,down judge,field judge,estadistico\n"
        "5,Campo 1,12:00,Nómadas U8,Lobos U8,Cam,Alan,Persona nueva,Allison\n"
    ).encode("utf-8")

    result = parse_referee_schedule_file(
        "rol-arbitros.csv", content, GAMES, REFEREES
    )

    proposal = result["proposals"][0]
    assert proposal["field_number"] == 1
    assert [official["position"] for official in proposal["officials"]] == [
        "referee", "down_judge", "statistician"
    ]
    assert result["matched"] == 1
    assert "Persona nueva" in result["warnings"][0]


def test_structured_referee_names_preserve_zoe_accents_and_reject_shared_aliases():
    people = [{"id":11,"name":"Zoe García","aka":"Zoe","display_name":"Zoe"},
              {"id":12,"name":"Zoé Pérez","aka":"Zoé","display_name":"Zoé"}]
    source = 'jornada,campo,hora,local,visitante,referee,down judge\n5,1,12:00,Nomadas U8,Lobos U8,Zoe,Zoé\n'.encode()
    for candidates in [people,list(reversed(people))]:
        result = parse_referee_schedule_file('rol.csv',source,GAMES,candidates)
        assert result['warnings']==[]
        assert [official['user_id'] for official in result['proposals'][0]['officials']]==[11,12]
    duplicate = [{**person,'aka':'Zoe','display_name':'Zoe'} for person in people]
    ambiguous = parse_referee_schedule_file('rol.csv',source,GAMES,duplicate)
    assert ambiguous['proposals'][0]['officials']==[]
    assert 'única' in ambiguous['warnings'][0]


def test_structured_schedule_does_not_guess_an_unregistered_game():
    content = (
        "jornada,campo,hora,local,visitante,referee,down judge\n"
        "5,1,12:00,Equipo inexistente,Lobos U8,Cam,Alan\n"
    ).encode("utf-8")

    result = parse_referee_schedule_file("rol.csv", content, GAMES, REFEREES)

    assert result["proposals"] == []
    assert result["unmatched"] == 1
    assert "no se encontro un partido seguro" in result["warnings"][0]


@pytest.mark.parametrize("home_label,away_label,home_name,away_name,home_category,away_category,branch", [
    ("Sel Ags U12 Fem", "Eagles Flag Fem 14", "Sel Ags U12 Fem", "EaglesFlag", "u12", "u14", "femenil"),
    ("Eagles Flag Fem 14", "Lobos Pla Fem 14", "EaglesFlag", "Lobos Pla", "u14", "u14", "femenil"),
    ("ECA Fem 14", "Eagles Flag Fem 14", "ECA", "EaglesFlag", "u14", "u14", "femenil"),
    ("Rancheras Flag Fem Libre", "Sel Ags U14 Fem Libre", "Rancheras Flag", "SelAgsFemU14", "libre", "libre", "femenil"),
    ("Lobos Inf Var Libre", "Nómadas Var Libre", "Lobos Kids", "Nómadas", "libre", "libre", "varonil"),
    ("Sel Ags U14 Fem Libre", "Storms Fem Libre", "SelAgsFemU14", "Storms", "libre", "libre", "femenil"),
    ("Lobos Inf Var Libre", "Vikingos Var Libre", "Lobos Kids", "Vikingos", "libre", "libre", "varonil"),
])
def test_referee_import_uses_the_registered_schedule_identity(
    home_label, away_label, home_name, away_name, home_category, away_category, branch,
):
    game = {**GAMES[0], "week": 3,
            "home_team": {"id": 1, "name": home_name, "branch": branch, "category": home_category},
            "away_team": {"id": 2, "name": away_name, "branch": branch, "category": away_category}}
    content = f"jornada,campo,hora,local,visitante,referee\n3,1,12:00,{home_label},{away_label},Cam\n".encode()
    result = parse_referee_schedule_file("rol.csv", content, [game], REFEREES)
    assert result["matched"] == 1
    assert result["warnings"] == []
    assert result["proposals"][0]["game_id"] == game["id"]


def test_referee_schedule_repeated_teams_use_slot_without_identity_ambiguity():
    content = b"jornada,campo,hora,local,visitante,referee\n5,2,13:00,Lobos U8,Nomadas U8,Cam\n"
    repeat = {**GAMES[0], "id": 82, "field_number": 2, "start_time": "13:00"}
    result = parse_referee_schedule_file("rol.csv", content, [*GAMES, repeat], REFEREES)
    assert result["proposals"][0]["game_id"] == 82
    ambiguous = {**repeat, "id": 83}
    result = parse_referee_schedule_file("rol.csv", content, [repeat, ambiguous], REFEREES)
    assert result["proposals"] == []


def test_referee_schedule_does_not_substitute_another_category_or_team():
    content = b"jornada,campo,hora,local,visitante,referee\n5,1,12:00,Nomadas U14,Lobos U8,Cam\n"
    assert parse_referee_schedule_file("rol.csv", content, GAMES, REFEREES)["matched"] == 0
    content = b"jornada,campo,hora,local,visitante,referee\n5,1,12:00,ECA U8,Lobos U8,Cam\n"
    assert parse_referee_schedule_file("rol.csv", content, GAMES, REFEREES)["matched"] == 0
