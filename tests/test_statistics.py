"""Integration tests for complete-week Excel imports and leaderboards."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from database import get_connection
from main import app


client = TestClient(app)
HEADERS = [
    "rama", "categoria", "equipo", "numero", "puntos", "recepciones",
    "intercepciones", "capturas", "tacleadas", "pases_completos",
    "pases_lanzados"
]


@pytest.fixture(autouse=True)
def clean_database():
    connection = get_connection()
    connection.execute("DELETE FROM player_week_stats")
    connection.execute("DELETE FROM games")
    connection.execute("DELETE FROM team_players")
    connection.execute("DELETE FROM players")
    connection.execute("DELETE FROM teams")
    connection.commit()
    connection.close()


def create_team(name, branch="varonil", category="libre"):
    return client.post(
        "/api/teams",
        json={"name": name, "branch": branch, "category": category}
    ).json()["id"]


def create_player(team_id, name, curp, number):
    return client.post(
        f"/api/teams/{team_id}/players",
        json={"name": name, "curp": curp, "age": 25, "jersey_number": number}
    ).json()["id"]


def create_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def create_official_workbook():
    """Build the horizontal event layout used by the league's real file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Estadísticas"
    week_one = workbook.create_sheet("Wk 1")
    week_one.append(["Equipo", "Categoría", "Estadística", 1, 2, 3, 4])
    labels = [
        "Intentos Pase", "Completos Pase", "Para % pases", "Puntos Pase",
        "6 Puntos", "2 Puntos", "1 Puntos", "Intercepciones Pase", "TD",
        "Conv 1", "Conv 2", "Sacks", "Tacleo", "Intercepciones Def",
        "Asistencia"
    ]
    tigres_events = {
        "Intentos Pase": [83], "Completos Pase": [10, 10],
        "Para % pases": [83, 83], "6 Puntos": [83], "TD": [10],
        "Sacks": [22, 22], "Tacleo": [22, 10],
        "Intercepciones Def": [22]
    }
    for label in labels:
        week_one.append(
            ["Tigres", "Var Libre", label, *tigres_events.get(label, [])]
        )
    for label in labels:
        week_one.append(["Ravens", "Var Libre", label])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def create_doubleheader_workbook(repeated_opponent=False):
    """Create two games where the same team plays twice in one jornada."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wk 5"
    sheet.append(["Equipo", "Categoría", "Estadística", 1, 2, 3, 4])
    labels = [
        "Intentos Pase", "Completos Pase", "Para % pases", "Puntos Pase",
        "6 Puntos", "2 Puntos", "1 Puntos", "Intercepciones Pase", "TD",
        "Conv 1", "Conv 2", "Sacks", "Tacleo", "Intercepciones Def",
        "Asistencia"
    ]
    second_game = (
        (("Kittys", 1, True), ("Diablos", 2, False))
        if repeated_opponent
        else (("Diablos", 2, True), ("APXGP", 3, False))
    )
    for team, number, points in (
        ("Kittys", 1, True), ("Diablos", 2, False), *second_game
    ):
        for label in labels:
            events = [number] if label == "Asistencia" or (points and label == "TD") else []
            sheet.append([team, "Mix Libre", label, *events])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def upload_week(week, rows, filename="estadisticas.xlsx"):
    return client.post(
        f"/api/weeks/{week}/player-stats/import",
        files={"file": (filename, create_workbook(rows))}
    )


def create_rosters():
    tigres = create_team("Tigres")
    ravens = create_team("Ravens")
    panteras = create_team("Panteras", "femenil", "u18")
    bruno = create_player(tigres, "Bruno Diaz", "DIBB961215HASXXX00", 83)
    jose = create_player(ravens, "Jose Perez", "PERJ990201HASXX002", 12)
    ana = create_player(panteras, "Ana Ruiz", "RUIA010101MASXX003", 7)
    return bruno, jose, ana


def test_imports_a_complete_week_and_resolves_names_from_rosters():
    bruno_id, _, _ = create_rosters()
    response = upload_week(2, [
        ["varonil", "libre", "Tigres", 83, 12, 4, 1, 0, 3, 18, 24],
        ["varonil", "libre", "Ravens", 12, 6, 2, 2, 1, 5, 11, 20],
        ["femenil", "u18", "Panteras", 7, 18, 5, 0, 0, 2, 0, 0]
    ])
    assert response.status_code == 200
    assert response.json()["imported"] == 3
    assert {row["player_name"] for row in response.json()["rows"]} == {
        "Bruno Diaz", "Jose Perez", "Ana Ruiz"
    }

    week = client.get("/api/weeks/2/player-stats")
    assert week.status_code == 200
    assert week.json()[0]["week"] == 2

    totals = client.get(f"/api/players/{bruno_id}/stats").json()
    assert totals["weeks"] == 1
    assert totals["points"] == 12
    assert totals["completion_percentage"] == 75.0


def test_statistics_import_resolves_u12_abbreviated_team_names():
    diablos = create_team("Diablos del Sol Amarillo", "femenil", "u12")
    player = create_player(
        diablos, "U12 Player", "U12X010101HASXX004", 4
    )

    response = upload_week(1, [[
        "femenil", "u12", "Diablos Ama", 4,
        6, 1, 0, 0, 2, 3, 5
    ]])

    assert response.status_code == 200
    assert response.json()["rows"][0]["branch"] == "mixto"
    totals = client.get(f"/api/players/{player}/stats").json()
    assert totals["points"] == 6


def test_imports_the_official_horizontal_workbook_format():
    tigres = create_team("Tigres")
    quarterback = create_player(tigres, "Quarterback", "QBXX010101HASXX001", 83)
    create_player(tigres, "Receiver", "RECX010101HASXX002", 10)
    create_player(tigres, "Defender", "DEFN010101HASXX003", 22)
    ravens = create_team("Ravens")

    response = client.post(
        "/api/statistics/import",
        files={"file": ("Stats ALL.xlsx", create_official_workbook())}
    )
    assert response.status_code == 200
    assert response.json()["weeks"] == [1]
    assert response.json()["imported"] == 3
    assert response.json()["games"] == 1

    totals = client.get(f"/api/players/{quarterback}/stats").json()
    assert totals["passes_attempted"] == 3
    assert totals["passes_completed"] == 2
    assert totals["points"] == 0
    leaders = client.get(
        "/api/statistics/leaderboards?branch=varonil&category=libre"
    ).json()
    assert leaders["receptions"][0]["player_name"] == "Receiver"
    assert leaders["points"][0]["value"] == 6
    assert leaders["sacks"][0]["value"] == 2
    games = client.get("/api/games").json()
    assert games[0]["home_team"]["id"] == tigres
    assert games[0]["away_team"]["id"] == ravens
    assert games[0]["home_score"] == 6
    assert games[0]["away_score"] == 0


def test_doubleheader_statistics_remain_scoped_to_each_game():
    """A team playing twice must expose the correct rows in both details."""
    kittys = create_team("Kittys", "mixto", "libre")
    diablos = create_team("Diablos", "mixto", "libre")
    apxgp = create_team("APXGP", "mixto", "libre")
    create_player(kittys, "Kittys One", "KITX010101HASXX001", 1)
    create_player(diablos, "Diablos Two", "DIAX010101HASXX002", 2)
    create_player(apxgp, "APXGP Three", "APXX010101HASXX003", 3)

    response = client.post(
        "/api/statistics/import",
        files={"file": ("doubleheader.xlsx", create_doubleheader_workbook())}
    )
    assert response.status_code == 200

    connection = get_connection()
    games = connection.execute(
        "SELECT id, home_team_id, away_team_id FROM games ORDER BY id"
    ).fetchall()
    first_teams = {
        row["team_id"] for row in connection.execute(
            "SELECT DISTINCT team_id FROM player_week_stats WHERE game_id = %s",
            (games[0]["id"],)
        ).fetchall()
    }
    second_teams = {
        row["team_id"] for row in connection.execute(
            "SELECT DISTINCT team_id FROM player_week_stats WHERE game_id = %s",
            (games[1]["id"],)
        ).fetchall()
    }
    connection.close()
    assert first_teams == {kittys, diablos}
    assert second_teams == {diablos, apxgp}


def test_repeated_matchup_creates_two_distinct_game_snapshots():
    """Two games between the same teams cannot share one statistics identity."""
    kittys = create_team("Kittys", "mixto", "libre")
    diablos = create_team("Diablos", "mixto", "libre")
    create_player(kittys, "Kittys One", "KITX010101HASXX001", 1)
    create_player(diablos, "Diablos Two", "DIAX010101HASXX002", 2)

    response = client.post(
        "/api/statistics/import",
        files={"file": (
            "repeated-matchup.xlsx",
            create_doubleheader_workbook(repeated_opponent=True)
        )}
    )
    assert response.status_code == 200

    connection = get_connection()
    games = connection.execute("SELECT id FROM games ORDER BY id").fetchall()
    game_ids = {
        row["game_id"] for row in connection.execute(
            "SELECT DISTINCT game_id FROM player_week_stats"
        ).fetchall()
    }
    connection.close()
    assert len(games) == 2
    assert game_ids == {games[0]["id"], games[1]["id"]}


def test_team_identity_includes_branch_and_category():
    first = create_team("Tigres", "varonil", "libre")
    second = create_team("Tigres", "varonil", "u18")
    create_player(first, "Adulto", "AAAA010101HASXX001", 10)
    create_player(second, "Juvenil", "BBBB010101HASXX002", 10)

    response = upload_week(1, [
        ["varonil", "libre", "Tigres", 10, 6, 1, 0, 0, 0, 1, 2],
        ["varonil", "u18", "Tigres", 10, 12, 2, 0, 0, 0, 2, 3]
    ])
    assert response.status_code == 200
    assert [row["player_name"] for row in response.json()["rows"]] == [
        "Adulto", "Juvenil"
    ]


def test_all_unknown_numbers_preserve_the_week_snapshot():
    create_rosters()
    assert upload_week(1, [
        ["varonil", "libre", "Tigres", 83, 12, 3, 0, 0, 1, 5, 10]
    ]).status_code == 200
    invalid = upload_week(1, [
        ["varonil", "libre", "Tigres", 999, 9, 9, 9, 9, 9, 5, 10]
    ])
    assert invalid.status_code == 200
    assert invalid.json()["imported"] == 0
    assert invalid.json()["skipped_count"] == 1
    assert invalid.json()["skipped"][0]["jersey_number"] == 999
    assert invalid.json()["preserved_weeks"] == [1]
    saved = client.get("/api/weeks/1/player-stats").json()
    assert len(saved) == 1
    assert saved[0]["points"] == 12


def test_unknown_numbers_are_skipped_only_in_their_own_team():
    create_rosters()
    response = upload_week(1, [
        ["varonil", "libre", "Tigres", 83, 12, 3, 0, 0, 1, 5, 10],
        ["varonil", "libre", "Tigres", 12, 9, 9, 0, 0, 0, 0, 0],
        ["varonil", "libre", "Ravens", 12, 6, 1, 0, 0, 0, 0, 0],
    ])
    assert response.status_code == 200
    data = response.json()
    assert data["imported"] == 2
    assert data["skipped_count"] == 1
    assert data["preserved_weeks"] == []
    assert data["skipped"][0]["team_name"] == "Tigres"
    assert data["skipped"][0]["week"] == 1
    assert data["skipped"][0]["jersey_number"] == 12
    assert len(client.get("/api/weeks/1/player-stats").json()) == 2


def test_official_import_reports_missing_numbers_and_keeps_valid_stats():
    tigres = create_team("Tigres")
    create_team("Ravens")
    create_player(tigres, "Quarterback", "QBXX010101HASXX001", 83)
    response = client.post(
        "/api/statistics/import",
        files={"file": ("stats.xlsx", create_official_workbook())},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["imported"] == 1
    assert data["games"] == 1
    assert {row["jersey_number"] for row in data["skipped"]} == {10, 22}
    assert data["skipped_count"] == 2
    assert len(client.get("/api/weeks/1/player-stats").json()) == 1


def test_unknown_team_still_rolls_back_after_skipped_number():
    create_rosters()
    upload_week(1, [
        ["varonil", "libre", "Tigres", 83, 12, 3, 0, 0, 1, 5, 10]
    ])
    response = upload_week(1, [
        ["varonil", "libre", "Tigres", 999, 0, 0, 0, 0, 0, 0, 0],
        ["varonil", "libre", "Unknown", 83, 6, 1, 0, 0, 0, 0, 0],
    ])
    assert response.status_code == 422
    assert client.get("/api/weeks/1/player-stats").json()[0]["points"] == 12


def test_reimport_updates_one_player_without_clearing_teammates_or_rivals():
    create_rosters()
    upload_week(1, [
        ["varonil", "libre", "Tigres", 83, 6, 1, 0, 0, 0, 1, 2],
        ["varonil", "libre", "Ravens", 12, 6, 1, 0, 0, 0, 1, 2]
    ])
    upload_week(1, [
        ["varonil", "libre", "Tigres", 83, 18, 3, 0, 0, 0, 2, 3]
    ])
    saved = client.get("/api/weeks/1/player-stats").json()
    assert len(saved) == 2
    assert {row["player_name"]: row["points"] for row in saved} == {
        "Bruno Diaz": 18, "Jose Perez": 6
    }


def test_import_accumulates_across_weeks_and_repeated_file_is_idempotent():
    player_id, _, _ = create_rosters()
    first = [["varonil", "libre", "Tigres", 83, 12, 2, 0, 0, 0, 7, 14]]
    second = [["varonil", "libre", "Tigres", 83, 18, 3, 0, 0, 0, 1, 1]]
    assert upload_week(1, first).status_code == 200
    assert upload_week(2, second).status_code == 200

    passing_url = "/api/statistics/leaderboards?branch=varonil&category=libre"
    before = client.get(passing_url).json()
    assert before["passing_qualification"]["minimum_attempts"] == 15
    assert before["completion_percentage"][0]["passes_attempted"] == 15
    assert before["completion_percentage"][0]["passes_completed"] == 8
    assert before["points"][0]["value"] == 30

    duplicate = upload_week(2, second)
    assert duplicate.status_code == 200
    assert duplicate.json()["updated"] == 1
    assert duplicate.json()["created"] == 0
    saved = client.get(f"/api/players/{player_id}/stats").json()
    assert saved["points"] == 30
    assert saved["passes_attempted"] == 15
    assert len(client.get("/api/weeks/2/player-stats").json()) == 1


def test_passers_under_fifteen_total_attempts_are_not_ranked():
    team = create_team("Lobos", "varonil", "libre")
    create_player(team, "Few Attempts", "FEWA010101HASXX001", 4)
    create_player(team, "Eligible", "ELIG010101HASXX002", 5)
    assert upload_week(1, [
        ["varonil", "libre", "Lobos", 4, 6, 0, 0, 0, 0, 14, 14],
        ["varonil", "libre", "Lobos", 5, 6, 0, 0, 0, 0, 10, 15],
    ]).status_code == 200
    result = client.get(
        "/api/statistics/leaderboards?branch=varonil&category=libre"
    ).json()
    assert [row["player_name"] for row in result["completion_percentage"]] == [
        "Eligible"
    ]
    assert result["passing_qualification"]["minimum_attempts"] == 15


def test_reimport_official_workbook_keeps_doubleheader_games_distinct():
    kittys = create_team("Kittys", "mixto", "libre")
    diablos = create_team("Diablos", "mixto", "libre")
    apxgp = create_team("APXGP", "mixto", "libre")
    create_player(kittys, "Kittys One", "KITX010101HASXX001", 1)
    create_player(diablos, "Diablos Two", "DIAX010101HASXX002", 2)
    create_player(apxgp, "APXGP Three", "APXX010101HASXX003", 3)
    excel = create_doubleheader_workbook()
    for _ in range(2):
        response = client.post(
            "/api/statistics/import", files={"file": ("week5.xlsx", excel)}
        )
        assert response.status_code == 200, response.text
    stats = client.get("/api/weeks/5/player-stats").json()
    assert len([row for row in stats if row["player_name"] == "Diablos Two"]) == 2
    assert len(stats) == 4


def test_import_rejects_non_excel_file_and_invalid_week():
    assert client.post(
        "/api/weeks/1/player-stats/import",
        files={"file": ("stats.csv", b"not excel", "text/csv")}
    ).status_code == 415
    assert client.post(
        "/api/weeks/0/player-stats/import",
        files={"file": ("stats.xlsx", create_workbook([]))}
    ).status_code == 422


def test_download_template_contains_roster_identity_columns():
    response = client.get("/api/statistics/import-template")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    headers = [cell.value for cell in next(workbook.active.iter_rows())]
    workbook.close()
    assert headers == HEADERS
    assert "nombre" not in headers


def test_leaderboards_limit_results_and_apply_passing_threshold():
    tigres = create_team("Tigres")
    rows = []
    for index in range(6):
        number = index + 1
        create_player(tigres, f"Player {number}", f"TEST{number:014d}", number)
        rows.append([
            "varonil", "libre", "Tigres", number,
            6 - index, 1, 0, 0, 0, 1, 2
        ])
    assert upload_week(1, rows).status_code == 200
    leaders = client.get(
        "/api/statistics/leaderboards?branch=varonil&category=libre"
    ).json()
    assert len(leaders["points"]) == 5
    assert leaders["passing_qualification"]["minimum_attempts"] == 15
    assert leaders["completion_percentage"] == []

    qualifying = rows[0].copy()
    qualifying[-2:] = [29, 38]
    assert upload_week(4, [qualifying]).status_code == 200
    leaders = client.get(
        "/api/statistics/leaderboards?branch=varonil&category=libre"
    ).json()
    assert leaders["passing_qualification"] == {
        "latest_week": 4, "minimum_attempts": 15
    }
    assert leaders["completion_percentage"][0]["player_name"] == "Player 1"
    assert "player_aka" in leaders["completion_percentage"][0]
    assert "profile_photo_url" in leaders["completion_percentage"][0]


def test_youth_leaderboards_combine_legacy_branches():
    connection = get_connection()
    legacy_team = connection.execute(
        """
        INSERT INTO teams (name, branch, category, status)
        VALUES ('Legacy Ravens', 'femenil', 'u10', 'active')
        RETURNING id
        """
    ).fetchone()["id"]
    connection.commit()
    connection.close()
    canonical_team = create_team("Nomadas", "mixto", "u10")
    create_player(legacy_team, "Legacy Leader", "LEGA010101MASXX001", 7)
    create_player(canonical_team, "Canonical Leader", "CANO010101HASXX002", 8)

    assert upload_week(1, [
        ["femenil", "u10", "Legacy Ravens", 7, 12, 2, 0, 0, 1, 0, 0],
        ["mixto", "u10", "Nomadas", 8, 6, 1, 0, 0, 1, 0, 0],
    ]).status_code == 200

    leaders = client.get(
        "/api/statistics/leaderboards?branch=varonil&category=u10"
    ).json()
    assert [leader["player_name"] for leader in leaders["points"]] == [
        "Legacy Leader", "Canonical Leader"
    ]


def test_leaderboards_share_places_and_keep_ties_at_cutoff():
    team = create_team('Tied Leaders')
    rows = []
    for index, (points, completed, attempted) in enumerate([
        (10, 1, 2), (10, 2, 4), (8, 3, 6),
        (6, 1, 3), (4, 1, 4), (4, 2, 8),
    ], 1):
        create_player(team, f'Tied Player {index}', f'TIED{index:014d}', index)
        rows.append(['varonil', 'libre', 'Tied Leaders', index,
                     points, 0, 0, 0, 0, completed * 8, attempted * 8])
    assert upload_week(1, rows).status_code == 200
    leaders = client.get(
        '/api/statistics/leaderboards?branch=varonil&category=libre'
    ).json()
    assert [row['rank'] for row in leaders['points']] == [1, 1, 3, 4, 5, 5]
    assert [row['rank'] for row in leaders['completion_percentage']] == [1, 2, 3, 4, 5]
    assert [row['passes_attempted'] for row in leaders['completion_percentage'][:3]] == [48, 32, 16]


def test_single_division_upload_preserves_other_divisions_in_same_week():
    """Uploading Fem U14 stats must not clear other divisions in that week."""
    fem_team = create_team("Cuernos", "femenil", "u14")
    var_team = create_team("Lobos", "varonil", "u14")
    create_player(fem_team, "Fem Player", "FEMU010101MASXX001", 4)
    create_player(var_team, "Var Player", "VARU010101HASXX001", 9)
    initial = upload_week(2, [
        ["femenil", "u14", "Cuernos", 4, 6, 1, 0, 0, 0, 0, 0],
        ["varonil", "u14", "Lobos", 9, 12, 1, 0, 0, 0, 0, 0],
    ])
    assert initial.status_code == 200

    response = client.post(
        "/api/weeks/2/player-stats/import-division",
        data={"branch": "femenil", "category": "u14"},
        files={"file": ("fem-u14.xlsx", create_workbook([
            ["femenil", "u14", "Cuernos", 4, 18, 3, 0, 0, 0, 0, 0]
        ]))},
    )
    assert response.status_code == 200
    assert response.json()["weeks"] == [2]
    stats = client.get("/api/weeks/2/player-stats").json()
    assert len(stats) == 2
    assert {row["team_name"]: row["points"] for row in stats} == {
        "Cuernos": 18,
        "Lobos": 12,
    }


def test_single_division_rejects_mixed_file_without_deleting_stats():
    create_rosters()
    assert upload_week(1, [
        ["femenil", "u18", "Panteras", 7, 6, 1, 0, 0, 0, 0, 0]
    ]).status_code == 200

    response = client.post(
        "/api/weeks/1/player-stats/import-division",
        data={"branch": "femenil", "category": "u18"},
        files={"file": ("mixed.xlsx", create_workbook([
            ["femenil", "u18", "Panteras", 7, 18, 0, 0, 0, 0, 0, 0],
            ["varonil", "libre", "Tigres", 83, 12, 0, 0, 0, 0, 0, 0],
        ]))},
    )
    assert response.status_code == 422
    saved = client.get("/api/weeks/1/player-stats").json()
    assert len(saved) == 1
    assert saved[0]["points"] == 6


def test_compact_bulk_upload_preserves_unmentioned_divisions():
    create_rosters()
    assert upload_week(2, [
        ["femenil", "u18", "Panteras", 7, 6, 1, 0, 0, 0, 0, 0],
        ["varonil", "libre", "Tigres", 83, 12, 2, 0, 0, 0, 0, 0],
    ]).status_code == 200

    response = upload_week(2, [
        ["varonil", "libre", "Tigres", 83, 24, 3, 0, 0, 0, 0, 0],
    ])
    assert response.status_code == 200
    saved = client.get("/api/weeks/2/player-stats").json()
    assert {row["team_name"]: row["points"] for row in saved} == {
        "Tigres": 24,
        "Panteras": 6,
    }


def test_official_workbook_preserves_unmentioned_categories():
    create_rosters()
    assert upload_week(1, [
        ["femenil", "u18", "Panteras", 7, 18, 1, 0, 0, 0, 0, 0]
    ]).status_code == 200

    response = client.post(
        "/api/statistics/import",
        files={"file": ("all.xlsx", create_official_workbook())},
    )
    assert response.status_code == 200
    saved = client.get("/api/weeks/1/player-stats").json()
    assert any(
        row["team_name"] == "Panteras" and row["points"] == 18
        for row in saved
    )


def test_official_workbook_does_not_overwrite_completed_game_scores():
    create_rosters()
    first = client.post(
        "/api/statistics/import",
        files={"file": ("all.xlsx", create_official_workbook())},
    )
    assert first.status_code == 200
    game_id = client.get("/api/games").json()[0]["id"]

    changed = client.patch(
        f"/api/games/{game_id}/score",
        json={"home_score": 21, "away_score": 0},
    )
    assert changed.status_code == 200

    imported = client.post(
        "/api/statistics/import",
        files={"file": ("all.xlsx", create_official_workbook())},
    )
    assert imported.status_code == 200
    game = client.get("/api/games").json()[0]
    assert game["home_score"] == 21
    assert game["away_score"] == 0


def create_single_sheet_u14_workbook():
    """Real-file shape: Hoja1 with 15 events per team, category U14 Fem."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    sheet.append(["Equipo", "Categoría", "Estadística", 1, 2, 3])
    labels = [
        "Intentos Pase", "Completos Pase", "Para % pases", "Puntos Pase",
        "6 Puntos", "2 Puntos", "1 Puntos", "Intercepciones Pase",
        "TD", "Conv 1", "Conv 2", "Sacks", "Tacleo",
        "Intercepciones Def", "Asistencia"
    ]
    for team, jersey, touchdowns in (
        ("Lobos", 7, 1), ("Ducks", 4, 0),
        ("Lobos", 7, 2), ("Pitbulls", 9, 0)
    ):
        for label in labels:
            events = (
                [jersey] * touchdowns if label == "TD"
                else [jersey] if label == "Asistencia"
                else [jersey, jersey] if label == "Intentos Pase"
                else []
            )
            sheet.append([team, "U14 Fem", label, *events])
    data = BytesIO()
    workbook.save(data)
    workbook.close()
    return data.getvalue()


def test_single_sheet_week1_fem_u14_links_existing_games_without_changing_scores():
    lobos = create_team("Lobos", "femenil", "u14")
    ducks = create_team("Ducks", "femenil", "u14")
    pitbulls = create_team("Pitbulls", "femenil", "u14")
    create_player(lobos, "Lobos QB", "LQBB010101MASXX004", 7)
    create_player(ducks, "Ducks Player", "DUPC010101MASXX005", 4)
    create_player(pitbulls, "Pitbulls Player", "PITC010101MASXX006", 9)
    connection = get_connection()
    first_id = connection.execute(
        """INSERT INTO games (home_team_id, away_team_id, week, status,
             home_score, away_score) VALUES (%s, %s, 1, 'completed', 21, 6)
             RETURNING id""", (lobos, ducks)
    ).fetchone()["id"]
    second_id = connection.execute(
        """INSERT INTO games (home_team_id, away_team_id, week, status,
             home_score, away_score) VALUES (%s, %s, 1, 'completed', 12, 13)
             RETURNING id""", (pitbulls, lobos)
    ).fetchone()["id"]
    connection.commit()
    connection.close()

    content = create_single_sheet_u14_workbook()
    for attempt in range(2):
        imported = client.post(
            "/api/weeks/1/player-stats/import-division",
            data={"branch": "femenil", "category": "u14"},
            files={"file": ("Week 1 U14 Femenil.xlsx", content)},
        )
        assert imported.status_code == 200, imported.text
        assert imported.json()["imported"] == 4
        assert imported.json()["updated"] == (0 if attempt == 0 else 4)

    saved = client.get("/api/weeks/1/player-stats").json()
    lobos_rows = [row for row in saved if row["player_name"] == "Lobos QB"]
    assert len(lobos_rows) == 2
    assert {row["game_id"] for row in lobos_rows} == {first_id, second_id}
    assert sum(row["points"] for row in lobos_rows) == 18
    assert sum(row["passes_attempted"] for row in lobos_rows) == 4

    connection = get_connection()
    games = connection.execute(
        "SELECT id, home_score, away_score FROM games ORDER BY id"
    ).fetchall()
    connection.close()
    assert [(g["id"], g["home_score"], g["away_score"]) for g in games] == [
        (first_id, 21, 6), (second_id, 12, 13)
    ]


def test_single_sheet_week2_fails_when_matches_only_exist_in_week1():
    lobos = create_team("Lobos", "femenil", "u14")
    ducks = create_team("Ducks", "femenil", "u14")
    pitbulls = create_team("Pitbulls", "femenil", "u14")
    create_player(lobos, "Lobos QB", "LQBB010101MASXX004", 7)
    create_player(ducks, "Ducks Player", "DUPC010101MASXX005", 4)
    create_player(pitbulls, "Pitbulls Player", "PITC010101MASXX006", 9)
    connection = get_connection()
    connection.execute(
        "INSERT INTO games (home_team_id, away_team_id, week) VALUES (%s, %s, 1)",
        (lobos, ducks)
    )
    connection.execute(
        "INSERT INTO games (home_team_id, away_team_id, week) VALUES (%s, %s, 1)",
        (pitbulls, lobos)
    )
    connection.commit()
    connection.close()
    imported = client.post(
        "/api/weeks/2/player-stats/import-division",
        data={"branch": "femenil", "category": "u14"},
        files={"file": ("Week 1 U14 Femenil.xlsx", create_single_sheet_u14_workbook())},
    )
    assert imported.status_code == 422, imported.text
    assert client.get("/api/weeks/2/player-stats").json() == []


def test_generic_team_name_variants_match_registered_teams():
    """Resolve real worksheet labels before validating jersey membership."""
    blackmamba = create_team("BlackMamba", "femenil", "u14")
    eagles = create_team("EaglesFlag", "femenil", "u14")
    diablos = create_team("DIABLOS DEL SOL", "femenil", "u14")
    create_player(blackmamba, "Black Player", "BLAK010101MASXX001", 1)
    create_player(eagles, "Eagles Player", "EAGL010101MASXX002", 2)
    create_player(diablos, "Diablos Player", "DIAB010101MASXX003", 3)

    response = upload_week(1, [
        ["femenil", "u14", "Black Mambas", 1, 6, 1, 0, 0, 0, 0, 0],
        ["femenil", "u14", "Eagles Flag", 2, 6, 1, 0, 0, 0, 0, 0],
        ["femenil", "u14", "Diablos", 3, 6, 1, 0, 0, 0, 0, 0],
    ])
    assert response.status_code == 200, response.text
    rows = client.get("/api/weeks/1/player-stats").json()
    assert {row["team_name"] for row in rows} == {
        "BlackMamba", "EaglesFlag", "DIABLOS DEL SOL"
    }


def test_team_name_variants_do_not_change_other_divisions():
    """An actual team named Diablos in another branch still resolves exactly."""
    original = create_team("Diablos", "varonil", "u14")
    create_player(original, "Different Diablos", "DIFD010101HASXX004", 7)
    response = upload_week(1, [
        ["varonil", "u14", "Diablos", 7, 6, 0, 0, 0, 0, 0, 0],
    ])
    assert response.status_code == 200, response.text
    assert client.get("/api/weeks/1/player-stats").json()[0]["team_name"] == "Diablos"


def test_spacing_and_plural_matching_applies_to_other_categories():
    """A novel spelling variant must work without a per-category alias."""
    team = create_team("BlackMamba", "femenil", "u16")
    create_player(team, "Athlete", "BLMB010101MASXX005", 11)
    response = upload_week(2, [
        ["femenil", "u16", "Black Mambas", 11, 6, 1, 0, 0, 0, 0, 0]
    ])
    assert response.status_code == 200, response.text
    assert client.get("/api/weeks/2/player-stats").json()[0]["team_name"] == "BlackMamba"


def test_ambiguous_team_name_stops_import_instead_of_guessing():
    """Similar teams in one division require an explicit choice."""
    first = create_team("BlackMamba", "femenil", "u14")
    second = create_team("Black Mambas", "femenil", "u14")
    create_player(first, "Player One", "PMBA010101MASXX006", 7)
    create_player(second, "Player Two", "PMBB010101MASXX007", 7)
    response = upload_week(1, [
        ["femenil", "u14", "Black Mamba", 7, 6, 0, 0, 0, 0, 0, 0]
    ])
    assert response.status_code == 422, response.text
    assert "Black Mamba" in response.json()["detail"]
    assert client.get("/api/weeks/1/player-stats").json() == []


def test_u12_same_name_in_two_branches_uses_registered_match():
    """The official opponent disambiguates Fem vs Mix U12 Nómadas."""
    eagles = create_team("Eagles Flag", "mixto", "u12")
    vikingos = create_team("Vikingos", "mixto", "u12")
    mictlan = create_team("Mictlán Bats", "mixto", "u12")
    rancheros = create_team("Rancheros", "mixto", "u12")
    nomadas_mixto = create_team("Nómadas", "mixto", "u12")
    # Legacy PROD has both names, but the public team creation API now
    # normalizes unified U12 requests to "mixto". Reproduce the historic
    # Femenil record directly, without loosening production validations.
    connection = get_connection()
    nomadas_fem = connection.execute(
        """
        INSERT INTO teams (name, branch, category, status)
        VALUES ('Nómadas', 'femenil', 'u12', 'active')
        RETURNING id
        """
    ).fetchone()["id"]
    connection.commit()
    connection.close()
    pitbulls = create_team("Pitbulls", "mixto", "u12")
    diablos = create_team("Diablos del Sol", "mixto", "u12")
    for team, jersey, label in [
        (eagles, 28, "Eagles"),
        (vikingos, 14, "Vikingos"),
        (mictlan, 10, "Mictlan"),
        (rancheros, 2, "Rancheros"),
        (nomadas_mixto, 16, "Wrong Nomadas"),
        (nomadas_fem, 15, "Right Nomadas"),
        (pitbulls, 13, "Pitbulls"),
        (diablos, 88, "Diablos"),
    ]:
        create_player(team, label, f"TESA010101HASXX{team % 1000:03d}", jersey)

    pairs = [
        (eagles, vikingos, 19, 41),
        (mictlan, rancheros, 0, 12),
        (nomadas_fem, pitbulls, 14, 0),
        (diablos, rancheros, 18, 13),
    ]
    connection = get_connection()
    for home, away, score_home, score_away in pairs:
        connection.execute(
            """
            INSERT INTO games
                (home_team_id, away_team_id, week, status, home_score, away_score)
            VALUES (%s, %s, 1, 'completed', %s, %s)
            """,
            (home, away, score_home, score_away)
        )
    connection.commit()
    connection.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    sheet.append(["Equipo", "Categoría", "Estadística", 1])
    labels = [
        "Intentos Pase", "Completos Pase", "Para % pases", "Puntos Pase",
        "6 Puntos", "2 Puntos", "1 Puntos", "Intercepciones Pase",
        "TD", "Conv 1", "Conv 2", "Sacks", "Tacleo",
        "Intercepciones Def", "Asistencia"
    ]
    for name, jersey in [
        ("Eagles Flag", 28), ("Vikingos", 14),
        ("Mictlan Bats", 10), ("Rancheros", 2),
        ("Nomadas", 15), ("Pitbulls", 13),
        ("Diablos", 88), ("Rancheros", 2),
    ]:
        for label in labels:
            sheet.append([
                name, "U12", label,
                jersey if label in ("Asistencia", "TD") else None,
            ])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    for attempt in range(2):
        response = client.post(
            "/api/weeks/1/player-stats/import-division",
            data={"branch": "mixto", "category": "u12"},
            files={"file": ("Week 1 U12.xlsx", stream.getvalue())},
        )
        assert response.status_code == 200, response.text
        assert response.json()["imported"] == 8
        assert response.json()["updated"] == (0 if attempt == 0 else 8)

    stats = client.get("/api/weeks/1/player-stats").json()
    assert len(stats) == 8
    assert len([row for row in stats if row["team_name"] == "Rancheros"]) == 2
    nomadas_rows = [row for row in stats if row["team_name"] == "Nómadas"]
    assert len(nomadas_rows) == 1
    assert nomadas_rows[0]["team_id"] == nomadas_fem
    assert nomadas_rows[0]["jersey_number"] == 15
    assert not any(row["team_id"] == nomadas_mixto for row in stats)

    connection = get_connection()
    scores = connection.execute(
        "SELECT home_score, away_score FROM games WHERE week = 1 ORDER BY id"
    ).fetchall()
    connection.close()
    assert [(game["home_score"], game["away_score"]) for game in scores] == [
        (19, 41), (0, 12), (14, 0), (18, 13)
    ]

def test_append_import_preserves_other_categories_and_is_idempotent():
    old_team = create_team('Existing U8', category='u8')
    create_player(old_team, 'Existing Player', 'AAAA000101HAGRRN01', 7)
    new_team = create_team('New Libre')
    create_player(new_team, 'New Player', 'BBBB000101HAGRRN02', 9)
    old_rows = [['mixto', 'u8', 'Existing U8', 7, 12, 1, 0, 0, 0, 0, 0]]
    new_rows = [['varonil', 'libre', 'New Libre', 9, 6, 2, 0, 0, 1, 0, 0]]
    path = '/api/weeks/1/player-stats/import'
    assert client.post(path, files={'file': ('old.xlsx', create_workbook(old_rows))}).status_code == 200
    for expected in (1, 0):
        response = client.post(path, data={'existing_records': 'preserve'}, files={'file': ('new.xlsx', create_workbook(new_rows))})
        assert response.status_code == 200
        assert response.json()['imported'] == expected
        assert response.json()['retained_count'] == 1 - expected
        rows = client.get('/api/weeks/1/player-stats').json()
        assert len(rows) == 2
        assert sorted(row['points'] for row in rows) == [6, 12]
    conflicting = [['varonil', 'libre', 'New Libre', 9, 99, 2, 0, 0, 1, 0, 0]]
    assert client.post(path, data={'existing_records': 'preserve'}, files={'file': ('changed.xlsx', create_workbook(conflicting))}).json()['imported'] == 0
    assert sorted(row['points'] for row in client.get('/api/weeks/1/player-stats').json()) == [6, 12]
    assert client.post(path, data={'existing_records': 'invalid'}, files={'file': ('bad.xlsx', create_workbook(new_rows))}).status_code == 422


def test_append_official_reupload_preserves_scores_and_does_not_duplicate():
    team = create_team('Tigres')
    create_team('Ravens')
    for number, name, curp in ((83, 'Passer', 'CCCC000101HAGRRN03'), (10, 'Receiver', 'DDDD000101HAGRRN04'), (22, 'Defender', 'EEEE000101HAGRRN05')):
        create_player(team, name, curp, number)
    payload = create_official_workbook()
    path = '/api/statistics/import'
    first = client.post(path, data={'existing_records': 'preserve'}, files={'file': ('official.xlsx', payload)})
    assert first.status_code == 200
    before = client.get('/api/weeks/1/player-stats').json()
    games_before = client.get('/api/games').json()
    second = client.post(path, data={'existing_records': 'preserve'}, files={'file': ('official.xlsx', payload)})
    assert second.status_code == 200
    assert second.json()['imported'] == 0
    assert second.json()['retained_count'] == len(before)
    assert client.get('/api/weeks/1/player-stats').json() == before
    assert client.get('/api/games').json() == games_before


def test_preserving_import_does_not_touch_unrelated_live_capture():
    home = create_team('Live Home', category='u8')
    away = create_team('Live Away', category='u8')
    create_player(home, 'Live Player', 'FFFF000101HAGRRN06', 1)
    game = client.post('/api/games', json={'home_team_id': home, 'away_team_id': away, 'field_number': 1, 'week': 1}).json()['id']
    with get_connection() as connection:
        connection.execute('INSERT INTO game_live_sessions (game_id) VALUES (%s)', (game,))
    other = create_team('Separate Libre')
    create_player(other, 'Separate Player', 'GGGG000101HAGRRN07', 9)
    rows = [['varonil', 'libre', 'Separate Libre', 9, 6, 0, 0, 0, 0, 0, 0]]
    endpoint = '/api/weeks/1/player-stats/import'
    response = client.post(endpoint, data={'existing_records': 'preserve'}, files={'file': ('new.xlsx', create_workbook(rows))})
    assert response.status_code == 200
    related = [['mixto', 'u8', 'Live Home', 1, 6, 0, 0, 0, 0, 0, 0]]
    assert client.post(endpoint, data={'existing_records': 'preserve'}, files={'file': ('live.xlsx', create_workbook(related))}).status_code == 422
    assert len(client.get('/api/weeks/1/player-stats').json()) == 1
    with get_connection() as connection:
        assert connection.execute('SELECT state FROM game_live_sessions WHERE game_id = %s', (game,)).fetchone()['state'] == 'live'


def test_original_official_team_labels_match_joined_registered_names():
    team = create_team('SelAgsFemU14', 'femenil', 'libre')
    create_team('Ravens', 'femenil', 'libre')
    for number, name, curp in ((83, 'Passer', 'HHHH000101HAGRRN08'), (10, 'Receiver', 'IIII000101HAGRRN09'), (22, 'Defender', 'JJJJ000101HAGRRN10')):
        create_player(team, name, curp, number)
    workbook = load_workbook(BytesIO(create_official_workbook()))
    for row in workbook['Wk 1'].iter_rows(min_row=2):
        if row[0].value == 'Tigres':
            row[0].value = 'Sel Ags U14'
        row[1].value = 'Libre Femenil'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = client.post('/api/statistics/import', data={'existing_records': 'preserve'}, files={'file': ('original.xlsx', output.getvalue())})
    assert response.status_code == 200
    rows = client.get('/api/weeks/1/player-stats').json()
    assert len(rows) == 3
    assert {row['team_id'] for row in rows} == {team}
    assert {row['team_name'] for row in rows} == {'SelAgsFemU14'}


def test_defensive_ranks_use_secondary_metrics_and_preserve_exact_ties():
    from repositories.statistics import _leader_score, _rank_leaders
    rows = [
        {"tackles": 10, "sacks": 3, "interceptions": 2},
        {"tackles": 10, "sacks": 3, "interceptions": 1},
        {"tackles": 10, "sacks": 2, "interceptions": 5},
        {"tackles": 10, "sacks": 2, "interceptions": 5},
    ]
    assert [r["rank"] for r in _rank_leaders(rows, lambda r: _leader_score(r, "tackles"), 5)] == [1, 2, 3, 3]
    assert _leader_score(rows[0], "sacks") == (3, 10)
    assert _leader_score(rows[0], "interceptions") == (2, 10)
