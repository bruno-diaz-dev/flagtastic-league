"""Recognize original roster layouts without rewriting uploaded workbooks."""
from io import BytesIO

import pytest
from openpyxl import Workbook

from services.roster_import import RosterImportError, parse_roster_file


def workbook_bytes(*, second_roster=False, missing_curp=False, invalid_number=False):
    workbook = Workbook()
    roster = workbook.active
    roster.title = 'Roster Flamingos'
    roster.append(['ROSTER FLAMINGOS'])
    roster.append(['TOCHITO BANDERA | TEMPORADA 2026–2027'])
    roster.append([])
    roster.append(['No.', 'Nombre de la jugadora', 'Identidad' if missing_curp else 'CURP', 'Edad', 'Jersey', 'Posición'])
    roster.append([1, 'Example Player', 'AAAA000101HAGRRN01', 26, 'invalid' if invalid_number else 77, 'QB'])
    notes = workbook.create_sheet('Notas')
    notes.append(['Revisar los datos antes de usar el roster'])
    workbook.active = 1  # Notes must never be mistaken for the roster.
    if second_roster:
        other = workbook.create_sheet('Otro equipo')
        other.append(['Nombre', 'CURP', 'Numero'])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_roster_after_titles_uses_jersey_not_row_number_and_ignores_notes():
    players = parse_roster_file('roster.xlsx', workbook_bytes())
    assert len(players) == 1
    assert players[0].jersey_number == 77
    assert players[0].name == 'Example Player'


def test_missing_required_columns_are_not_invented():
    with pytest.raises(RosterImportError, match='Faltan columnas'):
        parse_roster_file('roster.xlsx', workbook_bytes(missing_curp=True))


def test_multiple_roster_sheets_are_not_silently_combined():
    with pytest.raises(RosterImportError, match='varias hojas'):
        parse_roster_file('roster.xlsx', workbook_bytes(second_roster=True))


def test_validation_reports_original_excel_row():
    with pytest.raises(RosterImportError, match='Fila 5:'):
        parse_roster_file('roster.xlsx', workbook_bytes(invalid_number=True))
