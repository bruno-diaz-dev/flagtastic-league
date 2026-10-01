"""Parse structured referee schedules into administrator-reviewable assignments."""

from csv import reader, Sniffer
from datetime import datetime, time
from io import BytesIO, StringIO
import re
import unicodedata

from openpyxl import load_workbook

from services.referee_schedule_ocr import (
    GRID_FIVE_POSITION_ORDER,
    GRID_FOUR_POSITION_ORDER,
    _best_referee,
)
from services.team_matching import team_name_match_score


class RefereeScheduleImportError(Exception):
    """Raised when a structured referee schedule cannot be interpreted."""


MAX_ASSIGNMENTS = 500
BASE_HEADERS = {
    "week": {"jornada", "semana", "week"},
    "field_number": {"campo", "cancha", "field", "field number"},
    "start_time": {"hora", "horario", "time", "start time"},
    "home_team": {"local", "equipo local", "home", "home team"},
    "away_team": {"visitante", "equipo visitante", "away", "away team"},
}
OFFICIAL_HEADERS = {
    "referee": {"referee", "arbitro", "arbitro principal", "ref principal"},
    "down_judge": {"down judge", "dj"},
    "field_judge": {"field judge", "fj"},
    "side_judge": {"side judge", "sj"},
    "statistician": {"estadistico", "estadistica", "statistician", "stats"},
}


def parse_referee_schedule_file(filename, content, games, referees):
    """Read XLSX or CSV schedules without relying on OCR."""
    lower_name = (filename or "").lower()
    if lower_name.endswith(".xlsx"):
        rows = _read_workbook(content)
    elif lower_name.endswith(".csv"):
        rows = _read_csv(content)
    else:
        raise RefereeScheduleImportError("Se requiere un archivo XLSX o CSV")

    proposals = []
    warnings = []
    for row in rows:
        game = _match_game(row, games)
        if game is None:
            warnings.append(
                f"{row['source_row']}: no se encontro un partido seguro para "
                f"{row['home_team']} vs {row['away_team']}."
            )
            continue

        officials, unmatched_names = _match_officials(row["officials"], referees)
        if unmatched_names:
            warnings.append(
                f"{row['source_row']}: no se reconocieron los oficiales "
                f"{', '.join(unmatched_names)}."
            )
        proposals.append({
            "game_id": game["id"],
            "game_label": (
                f"J{game['week']}: {game['home_team']['name']} vs "
                f"{game['away_team']['name']}"
            ),
            "field_number": row["field_number"],
            "scheduled_time": row["start_time"].strftime("%H:%M"),
            "officials": officials,
            "source_text": row["source_text"],
        })

    if not rows:
        raise RefereeScheduleImportError(
            "El archivo no contiene asignaciones arbitrales reconocibles"
        )
    if len(rows) > MAX_ASSIGNMENTS:
        raise RefereeScheduleImportError(
            f"El archivo no puede contener mas de {MAX_ASSIGNMENTS} partidos"
        )
    return {
        "proposals": proposals,
        "raw_text": "\n".join(row["source_text"] for row in rows),
        "warnings": warnings,
        "matched": len(proposals),
        "unmatched": len(rows) - len(proposals),
        "source_format": "structured",
    }


def _read_workbook(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise RefereeScheduleImportError("No se pudo leer el archivo XLSX") from error

    rows = []
    try:
        for sheet in workbook.worksheets:
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            rows.extend(_read_matrix(matrix, sheet.title))
    finally:
        workbook.close()
    return rows


def _read_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise RefereeScheduleImportError("El CSV debe estar codificado en UTF-8") from error
    try:
        dialect = Sniffer().sniff(text[:4096], delimiters=",;")
    except Exception:
        dialect = "excel"
    matrix = [list(row) for row in reader(StringIO(text), dialect=dialect)]
    return _read_matrix(matrix, "CSV")


def _read_matrix(matrix, source_name):
    """Accept either one-row-per-game data or the league's two-row grid."""
    header_index, columns = _tabular_columns(matrix)
    if columns is not None:
        return _read_tabular_rows(matrix, header_index, columns, source_name)
    return _read_grid_rows(matrix, source_name)


def _tabular_columns(matrix):
    for index, row in enumerate(matrix[:20]):
        normalized = {_normalize(value): column for column, value in enumerate(row) if _text(value)}
        columns = {}
        for key, aliases in BASE_HEADERS.items():
            match = next((normalized[alias] for alias in aliases if alias in normalized), None)
            if match is None:
                break
            columns[key] = match
        else:
            official_columns = {
                position: next(
                    (normalized[alias] for alias in aliases if alias in normalized),
                    None,
                )
                for position, aliases in OFFICIAL_HEADERS.items()
            }
            if any(column is not None for column in official_columns.values()):
                columns["officials"] = official_columns
                return index, columns
    return None, None


def _read_tabular_rows(matrix, header_index, columns, source_name):
    rows = []
    for row_number, values in enumerate(matrix[header_index + 1:], start=header_index + 2):
        if not any(_text(value) for value in values):
            continue
        try:
            week = int(re.sub(r"\D", "", _cell(values, columns["week"])))
            field_number = int(re.sub(r"\D", "", _cell(values, columns["field_number"])))
        except ValueError as error:
            raise RefereeScheduleImportError(
                f"{source_name}!{row_number}: jornada o campo invalido"
            ) from error
        start_time = _time_value(_cell_value(values, columns["start_time"]))
        home_team = _cell(values, columns["home_team"])
        away_team = _cell(values, columns["away_team"])
        if week < 1 or field_number not in range(1, 9) or start_time is None:
            raise RefereeScheduleImportError(
                f"{source_name}!{row_number}: revisa jornada, campo y hora"
            )
        officials = [
            (position, _cell(values, column))
            for position, column in columns["officials"].items()
            if column is not None and _cell(values, column)
        ]
        rows.append(_schedule_row(
            source_name, row_number, week, field_number, start_time,
            home_team, away_team, officials,
        ))
    return rows


def _read_grid_rows(matrix, source_name):
    week = _week(source_name)
    hour_column = None
    field_columns = {}
    pending_games = {}
    rows = []

    for row_number, values in enumerate(matrix, start=1):
        texts = [_text(value) for value in values]
        detected_week = next((_week(value) for value in texts if _week(value)), None)
        if detected_week is not None:
            week = detected_week

        detected_hour = next(
            (index for index, value in enumerate(texts) if _normalize(value) == "hora"),
            None,
        )
        if detected_hour is not None:
            hour_column = detected_hour
            field_columns = {}
            for column, value in enumerate(texts):
                match = re.fullmatch(r"(?:campo|cancha)\s*([1-8])", _normalize(value))
                if match:
                    field_columns[column] = int(match.group(1))
            continue

        if week is None or hour_column is None or not field_columns:
            continue
        start_time = _time_value(_cell_value(values, hour_column))
        if start_time is not None:
            pending_games = {}
            for column, field_number in field_columns.items():
                home_team = _cell(values, column)
                away_team = _cell(values, column + 1)
                if home_team and away_team:
                    pending_games[column] = (
                        field_number, start_time, home_team, away_team, row_number
                    )
            continue

        # The league sheet prints slash-separated officials directly below
        # each pair of team cells. Empty slash slots retain their role index.
        for column, game_data in list(pending_games.items()):
            official_text = " ".join(
                value for value in (_cell(values, column), _cell(values, column + 1)) if value
            ).strip()
            if "/" not in official_text:
                continue
            field_number, scheduled_time, home_team, away_team, game_row = game_data
            segments = [segment.strip() for segment in official_text.split("/")]
            order = GRID_FIVE_POSITION_ORDER if len(segments) >= 5 else GRID_FOUR_POSITION_ORDER
            officials = [
                (position, segment)
                for position, segment in zip(order, segments)
                if segment
            ]
            rows.append(_schedule_row(
                source_name, game_row, week, field_number, scheduled_time,
                home_team, away_team, officials,
            ))
            del pending_games[column]
    return rows


def _schedule_row(source, row_number, week, field_number, start_time, home, away, officials):
    names = " / ".join(name for _position, name in officials)
    return {
        "source_row": f"{source}!{row_number}",
        "week": week,
        "field_number": field_number,
        "start_time": start_time,
        "home_team": home,
        "away_team": away,
        "officials": officials,
        "source_text": (
            f"{source}!{row_number}: {home} vs {away} | {names}"
        ),
    }


def _match_game(row, games):
    candidates = []
    for game in games:
        if game["week"] != row["week"]:
            continue
        direct = (
            team_name_match_score(row["home_team"], game["home_team"]["name"]),
            team_name_match_score(row["away_team"], game["away_team"]["name"]),
        )
        swapped = (
            team_name_match_score(row["home_team"], game["away_team"]["name"]),
            team_name_match_score(row["away_team"], game["home_team"]["name"]),
        )
        team_scores = max(direct, swapped, key=sum)
        if min(team_scores) < 0.70 or sum(team_scores) < 1.50:
            continue
        field_bonus = 0.35 if game.get("field_number") == row["field_number"] else 0
        time_bonus = 0.35 if _same_time(game.get("start_time"), row["start_time"]) else 0
        candidates.append((sum(team_scores) + field_bonus + time_bonus, game))
    candidates.sort(key=lambda candidate: candidate[0], reverse=True)
    if not candidates:
        return None
    if len(candidates) > 1 and candidates[0][0] - candidates[1][0] < 0.08:
        return None
    return candidates[0][1]


def _match_officials(official_slots, referees):
    officials = []
    unmatched = []
    used_users = set()
    for position, printed_name in official_slots:
        matched = _best_referee(printed_name, referees)
        if matched is None:
            unmatched.append(printed_name)
            continue
        referee, display_name = matched
        if referee["id"] in used_users:
            unmatched.append(printed_name)
            continue
        used_users.add(referee["id"])
        officials.append({
            "user_id": referee["id"],
            "name": display_name,
            "position": position,
        })
    return officials, unmatched


def _same_time(left, right):
    left_time = _time_value(left)
    right_time = _time_value(right)
    return left_time is not None and right_time is not None and left_time == right_time


def _time_value(value):
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)) and 0 <= value < 1:
        minutes = round(value * 24 * 60)
        return time((minutes // 60) % 24, minutes % 60)
    text = _text(value).lower().replace("a. m.", "am").replace("p. m.", "pm")
    text = text.replace("a.m.", "am").replace("p.m.", "pm")
    for pattern in ("%H:%M", "%H.%M", "%I:%M %p", "%I %p"):
        try:
            return datetime.strptime(text.upper(), pattern).time()
        except ValueError:
            continue
    return None


def _week(value):
    match = re.search(r"(?:semana|jornada|week)\s*#?\s*(\d{1,2})", _normalize(value))
    return int(match.group(1)) if match else None


def _normalize(value):
    text = unicodedata.normalize("NFKD", _text(value))
    return " ".join(
        "".join(character for character in text if not unicodedata.combining(character))
        .casefold()
        .replace("_", " ")
        .split()
    )


def _text(value):
    return str(value).strip() if value is not None else ""


def _cell(values, index):
    return _text(_cell_value(values, index))


def _cell_value(values, index):
    return values[index] if index < len(values) else None
