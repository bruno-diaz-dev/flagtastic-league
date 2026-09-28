"""Parse league game schedules from the supported spreadsheet formats."""

from csv import DictReader, Sniffer
from datetime import datetime, time
from difflib import SequenceMatcher
from io import BytesIO, StringIO
import re
import unicodedata

from openpyxl import load_workbook
from PIL import Image, ImageEnhance, ImageOps, UnidentifiedImageError
import pytesseract

from services.referee_schedule_ocr import _configure_windows_tesseract


class GameScheduleImportError(Exception):
    """Raised when a schedule file cannot be interpreted safely."""


MAX_SCHEDULE_GAMES = 500
MAX_SCHEDULE_FIELDS = 6
CSV_HEADERS = {
    "week": {"week", "jornada", "semana"},
    "field_number": {"field", "field_number", "campo", "cancha"},
    "start_time": {"time", "start_time", "hora", "horario"},
    "home_team": {"home", "home_team", "local", "equipo_local"},
    "away_team": {"away", "away_team", "visitante", "equipo_visitante"},
}


def parse_game_schedule_file(filename, content, teams):
    """Return reviewable game proposals without modifying league data."""
    lower_name = (filename or "").lower()
    if lower_name.endswith(".xlsx"):
        rows = _read_calendar_workbook(content)
    elif lower_name.endswith(".csv"):
        rows = _read_csv(content)
    else:
        raise GameScheduleImportError("Se requiere un archivo XLSX o CSV")

    proposals = [_match_row(row, teams) for row in rows]
    if not proposals:
        raise GameScheduleImportError("El archivo no contiene partidos reconocibles")
    if len(proposals) > MAX_SCHEDULE_GAMES:
        raise GameScheduleImportError(
            f"El archivo no puede contener mas de {MAX_SCHEDULE_GAMES} partidos"
        )
    return {
        "kind": "games",
        "proposals": proposals,
        "matched": sum(proposal["ready"] for proposal in proposals),
        "unmatched": sum(not proposal["ready"] for proposal in proposals),
    }


def parse_game_schedule_image(content, teams):
    """Read the league's visual game grid into reviewable game proposals."""
    _configure_windows_tesseract()
    try:
        image = Image.open(BytesIO(content))
        image.verify()
        image = Image.open(BytesIO(content)).convert("L")
    except (UnidentifiedImageError, OSError) as error:
        raise GameScheduleImportError("La imagen no es valida") from error

    image = ImageOps.autocontrast(image.resize((image.width * 2, image.height * 2)))
    image = ImageEnhance.Contrast(image).enhance(1.5)
    try:
        raw_text = pytesseract.image_to_string(image, config="--psm 6")
        data = pytesseract.image_to_data(
            image, config="--psm 6", output_type=pytesseract.Output.DICT
        )
    except pytesseract.TesseractNotFoundError as error:
        raise GameScheduleImportError(
            "El servidor no tiene instalado el lector OCR"
        ) from error

    week = _week(raw_text)
    field_numbers = {
        int(number) for number in re.findall(rf"campo\s*([1-{MAX_SCHEDULE_FIELDS}])", _normalize(raw_text))
    }
    if week is None or not field_numbers:
        raise GameScheduleImportError(
            "No se reconocieron la jornada y los encabezados de campo"
        )
    field_count = max(field_numbers)
    time_rows = []
    for index, value in enumerate(data["text"]):
        parsed_time = _time_value(value)
        if parsed_time is None:
            continue
        center_y = data["top"][index] + data["height"][index] / 2
        if not any(abs(center_y - existing[0]) < image.height * 0.015 for existing in time_rows):
            time_rows.append((center_y, parsed_time))
    time_rows.sort(key=lambda item: item[0])
    if not time_rows:
        raise GameScheduleImportError("No se reconocieron horarios en la imagen")

    time_column_width = image.width * 0.027
    field_width = (image.width - time_column_width) / field_count
    proposals = []
    for row_index, (center_y, scheduled_time) in enumerate(time_rows):
        previous_y = time_rows[row_index - 1][0] if row_index else center_y - image.height * 0.04
        next_y = time_rows[row_index + 1][0] if row_index + 1 < len(time_rows) else center_y + image.height * 0.04
        top = max(0, round((previous_y + center_y) / 2))
        bottom = min(image.height, round((center_y + next_y) / 2))
        for field_index in range(field_count):
            left = time_column_width + field_index * field_width
            middle = left + field_width / 2
            right = left + field_width
            home_text = _ocr_team_cell(image, left, middle, top, bottom)
            away_text = _ocr_team_cell(image, middle, right, top, bottom)
            if not home_text or not away_text:
                continue
            proposals.append(_match_row({
                "source_row": f"Imagen, fila {row_index + 1}",
                "week": week,
                "field_number": field_index + 1,
                "start_time": scheduled_time,
                "home_team": home_text,
                "away_team": away_text,
            }, teams))

    if not proposals:
        raise GameScheduleImportError("No se reconocieron partidos en la imagen")
    return {
        "kind": "games",
        "proposals": proposals,
        "matched": sum(proposal["ready"] for proposal in proposals),
        "unmatched": sum(not proposal["ready"] for proposal in proposals),
        "raw_text": raw_text.strip(),
    }


def parse_game_schedule_ocr_words(
    image_width,
    image_height,
    words,
    teams,
    recognized_text=None,
):
    """Build schedule proposals from browser-side OCR words and coordinates."""
    normalized_words = [
        {
            "text": _text(word["text"]),
            "left": int(word["left"]),
            "top": int(word["top"]),
            "width": int(word["width"]),
            "height": int(word["height"]),
        }
        for word in words
        if _text(word.get("text"))
    ]
    word_text = " ".join(word["text"] for word in normalized_words)
    raw_text = _text(recognized_text) or word_text
    week = _week(raw_text) or _ocr_week_from_words(
        normalized_words,
        image_height,
    )
    if week is None:
        raise GameScheduleImportError("No se reconoció la jornada en la imagen")

    time_rows = _ocr_time_rows_from_words(
        normalized_words,
        image_width,
        image_height,
    )
    if not time_rows:
        raise GameScheduleImportError("No se reconocieron horarios en la imagen")

    field_numbers = {
        int(number)
        for number in re.findall(rf"campo\s*([1-{MAX_SCHEDULE_FIELDS}])", _normalize(raw_text))
    }
    if field_numbers:
        field_count = max(field_numbers)
    else:
        field_count = _ocr_field_count_from_words(
            normalized_words,
            image_width,
            image_height,
        )
    if field_count is None:
        field_count = _infer_field_count_from_team_matches(
            normalized_words,
            image_width,
            image_height,
            time_rows,
            teams,
        )
    if field_count is None:
        raise GameScheduleImportError(
            "No se pudo determinar la cantidad de campos en la imagen"
        )

    time_column_width = image_width * 0.027
    field_width = (image_width - time_column_width) / field_count
    proposals = []
    for row_index, (_center_y, scheduled_time) in enumerate(time_rows):
        top, bottom = _ocr_row_bounds(time_rows, row_index, image_height)
        for field_index in range(field_count):
            left = time_column_width + field_index * field_width
            middle = left + field_width / 2
            right = left + field_width
            home_text = _words_in_cell(
                normalized_words, left, middle, top, bottom
            )
            away_text = _words_in_cell(
                normalized_words, middle, right, top, bottom
            )
            if not home_text or not away_text:
                continue
            proposals.append(_match_row({
                "source_row": f"Imagen, fila {row_index + 1}",
                "week": week,
                "field_number": field_index + 1,
                "start_time": scheduled_time,
                "home_team": home_text,
                "away_team": away_text,
            }, teams))

    if not proposals:
        raise GameScheduleImportError("No se reconocieron partidos en la imagen")
    if len(proposals) > MAX_SCHEDULE_GAMES:
        raise GameScheduleImportError(
            f"La imagen no puede contener mas de {MAX_SCHEDULE_GAMES} partidos"
        )
    return {
        "kind": "games",
        "proposals": proposals,
        "matched": sum(proposal["ready"] for proposal in proposals),
        "unmatched": sum(not proposal["ready"] for proposal in proposals),
        "raw_text": raw_text.strip(),
    }


def _ocr_time_value(value):
    """Parse common OCR variants of HH:MM."""
    text = _text(value).strip()
    if not text:
        return None

    cleaned = (
        text.upper()
        .replace("O", "0")
        .replace("I", "1")
        .replace("L", "1")
        .replace(".", ":")
        .replace(";", ":")
        .replace(",", ":")
    )
    cleaned = re.sub(r"\s+", "", cleaned)

    match = re.fullmatch(r"(\d{1,2}):?(\d{2})", cleaned)
    if match:
        hour = int(match.group(1))
        minute = int(match.group(2))
        if 0 <= hour <= 23 and 0 <= minute <= 59:
            return time(hour, minute)
    return None


def _ocr_time_rows_from_words(words, image_width, image_height):
    """Recover schedule times from the narrow left-hand time column."""
    time_column_limit = image_width * 0.055
    candidates = [
        word for word in words
        if word["left"] + word["width"] / 2 <= time_column_limit
    ]
    candidates.sort(key=lambda word: (word["top"], word["left"]))

    tolerance = max(3, image_height * 0.012)
    bands = []
    for word in candidates:
        center_y = word["top"] + word["height"] / 2
        band = next(
            (
                existing for existing in bands
                if abs(existing["center_y"] - center_y) <= tolerance
            ),
            None,
        )
        if band is None:
            bands.append({"center_y": center_y, "words": [word]})
        else:
            band["words"].append(word)
            centers = [
                item["top"] + item["height"] / 2
                for item in band["words"]
            ]
            band["center_y"] = sum(centers) / len(centers)

    rows = []
    for band in bands:
        ordered = sorted(band["words"], key=lambda word: word["left"])
        variants = [
            "".join(word["text"] for word in ordered),
            " ".join(word["text"] for word in ordered),
            *[word["text"] for word in ordered],
        ]
        parsed = next(
            (_ocr_time_value(value) for value in variants if _ocr_time_value(value)),
            None,
        )
        if parsed is not None:
            rows.append((band["center_y"], parsed))

    # Some OCR engines position the time slightly outside the narrow column.
    # Fall back to any standalone OCR token only when the column yielded none.
    if not rows:
        for word in words:
            parsed = _ocr_time_value(word["text"])
            if parsed is None:
                continue
            center_y = word["top"] + word["height"] / 2
            if not any(abs(center_y - existing[0]) < tolerance for existing in rows):
                rows.append((center_y, parsed))

    rows.sort(key=lambda item: item[0])
    return rows


def _ocr_row_bounds(time_rows, row_index, image_height):
    center_y = time_rows[row_index][0]
    previous_y = (
        time_rows[row_index - 1][0]
        if row_index
        else center_y - image_height * 0.04
    )
    next_y = (
        time_rows[row_index + 1][0]
        if row_index + 1 < len(time_rows)
        else center_y + image_height * 0.04
    )
    return (
        max(0, (previous_y + center_y) / 2),
        min(image_height, (center_y + next_y) / 2),
    )


def _infer_field_count_from_team_matches(
    words,
    image_width,
    image_height,
    time_rows,
    teams,
):
    """Infer the grid width by choosing the layout that matches most teams."""
    time_column_width = image_width * 0.027
    best = None

    for field_count in range(1, MAX_SCHEDULE_FIELDS + 1):
        field_width = (image_width - time_column_width) / field_count
        ready_games = 0
        matched_sides = 0
        populated_cells = 0

        for row_index, _row in enumerate(time_rows):
            top, bottom = _ocr_row_bounds(time_rows, row_index, image_height)
            for field_index in range(field_count):
                left = time_column_width + field_index * field_width
                middle = left + field_width / 2
                right = left + field_width
                home_text = _words_in_cell(words, left, middle, top, bottom)
                away_text = _words_in_cell(words, middle, right, top, bottom)
                if not home_text or not away_text:
                    continue

                populated_cells += 1
                home = _match_team(home_text, teams)
                away = _match_team(away_text, teams)
                matched_sides += int(home is not None) + int(away is not None)
                if home is not None and away is not None and home["id"] != away["id"]:
                    ready_games += 1

        candidate = (ready_games, matched_sides, -populated_cells, field_count)
        if best is None or candidate > best:
            best = candidate

    if best is None or best[0] == 0:
        return None
    return best[3]


def _ocr_week_from_words(words, image_height):
    """Recover the week number when OCR misses the 'Semana' label."""
    top_limit = image_height * 0.10
    numeric_words = []
    for word in words:
        center_y = word["top"] + word["height"] / 2
        normalized = _normalize(word["text"])
        if center_y > top_limit or not re.fullmatch(r"\d{1,2}", normalized):
            continue
        value = int(normalized)
        if 1 <= value <= 30:
            numeric_words.append((center_y, word["left"], value))
    if not numeric_words:
        return None
    numeric_words.sort()
    return numeric_words[0][2]


def _ocr_field_count_from_words(words, image_width, image_height):
    """Recover field count from the numbered header row without its labels."""
    top = image_height * 0.08
    bottom = image_height * 0.28
    numbered = []
    for word in words:
        center_x = word["left"] + word["width"] / 2
        center_y = word["top"] + word["height"] / 2
        normalized = _normalize(word["text"])
        if (
            not (top <= center_y <= bottom)
            or center_x < image_width * 0.04
            or not re.fullmatch(rf"[1-{MAX_SCHEDULE_FIELDS}]", normalized)
        ):
            continue
        numbered.append((center_y, center_x, int(normalized)))

    if not numbered:
        return None

    # Header numbers share nearly the same baseline. Pick the densest band so
    # unrelated numbers elsewhere in the top area do not become field labels.
    tolerance = max(4, image_height * 0.025)
    best_band = []
    for anchor_y, _x, _value in numbered:
        band = [
            candidate for candidate in numbered
            if abs(candidate[0] - anchor_y) <= tolerance
        ]
        if len(band) > len(best_band):
            best_band = band

    values = sorted({value for _y, _x, value in best_band})
    if len(values) < 2 or values[0] != 1:
        return None

    maximum = values[-1]
    expected = list(range(1, maximum + 1))
    return maximum if values == expected else None


def _words_in_cell(words, left, right, top, bottom):
    selected = []
    for word in words:
        center_x = word["left"] + word["width"] / 2
        center_y = word["top"] + word["height"] / 2
        if left <= center_x < right and top <= center_y < bottom:
            selected.append(word)
    selected.sort(key=lambda word: (word["top"], word["left"]))
    return " ".join(word["text"] for word in selected).strip()


def _ocr_team_cell(image, left, right, top, bottom):
    crop = ImageOps.autocontrast(image.crop((round(left), top, round(right), bottom)))
    return pytesseract.image_to_string(crop, config="--psm 6").strip().replace("\n", " ")


def _read_calendar_workbook(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise GameScheduleImportError("No se pudo leer el archivo XLSX") from error

    rows = []
    try:
        for sheet in workbook.worksheets:
            week = None
            field_columns = {}
            hour_column = None
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = list(values)
                texts = [_text(value) for value in cells]
                week_value = next((value for value in texts if _week(value)), None)
                if week_value:
                    week = _week(week_value)
                    field_columns = {}
                    hour_column = None
                    continue

                detected_hour_column = next(
                    (index for index, value in enumerate(texts) if _normalize(value) == "hora"),
                    None,
                )
                if detected_hour_column is not None:
                    hour_column = detected_hour_column
                    field_columns = {}
                    for index, value in enumerate(texts):
                        match = re.fullmatch(rf"campo\s*([1-{MAX_SCHEDULE_FIELDS}])", _normalize(value))
                        if match and index + 1 < len(cells):
                            field_columns[index] = int(match.group(1))
                    continue

                if week is None or hour_column is None or not field_columns:
                    continue
                scheduled_time = _time_value(cells[hour_column])
                if scheduled_time is None:
                    continue
                for column, field_number in field_columns.items():
                    home = _text(cells[column])
                    away = _text(cells[column + 1])
                    if not home or not away or "copa aguascalientes" in _normalize(home):
                        continue
                    rows.append({
                        "source_row": f"{sheet.title}!{row_number}",
                        "week": week,
                        "field_number": field_number,
                        "start_time": scheduled_time,
                        "home_team": home,
                        "away_team": away,
                    })
    finally:
        workbook.close()
    return rows


def _read_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise GameScheduleImportError("El CSV debe estar codificado en UTF-8") from error
    try:
        dialect = Sniffer().sniff(text[:4096], delimiters=",;")
    except Exception:
        dialect = "excel"
    reader = DictReader(StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise GameScheduleImportError("El CSV no contiene encabezados")
    normalized = {_normalize(header): header for header in reader.fieldnames}
    mapped = {}
    for field, aliases in CSV_HEADERS.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match is None:
            raise GameScheduleImportError(
                "El CSV requiere jornada, campo, hora, local y visitante"
            )
        mapped[field] = match

    rows = []
    for row_number, row in enumerate(reader, start=2):
        if not any(_text(value) for value in row.values()):
            continue
        try:
            week = int(_text(row[mapped["week"]]))
            field_number = int(re.sub(r"\D", "", _text(row[mapped["field_number"]])))
        except (TypeError, ValueError) as error:
            raise GameScheduleImportError(
                f"Fila {row_number}: jornada o campo invalido"
            ) from error
        scheduled_time = _time_value(row[mapped["start_time"]])
        if week < 1 or field_number not in range(1, 9) or scheduled_time is None:
            raise GameScheduleImportError(f"Fila {row_number}: revisa jornada, campo y hora")
        rows.append({
            "source_row": str(row_number),
            "week": week,
            "field_number": field_number,
            "start_time": scheduled_time,
            "home_team": _text(row[mapped["home_team"]]),
            "away_team": _text(row[mapped["away_team"]]),
        })
    return rows


def _match_row(row, teams):
    home = _match_team(row["home_team"], teams)
    away = _match_team(row["away_team"], teams)
    ready = home is not None and away is not None and home["id"] != away["id"]
    return {
        **row,
        "start_time": row["start_time"].strftime("%H:%M"),
        "home_team_id": home["id"] if home else None,
        "home_match": home["name"] if home else None,
        "away_team_id": away["id"] if away else None,
        "away_match": away["name"] if away else None,
        "ready": ready,
    }


def _match_team(label, teams):
    target = _normalize(label)
    candidates = []
    for team in teams:
        score = max(
            SequenceMatcher(None, target, alias).ratio()
            for alias in _team_aliases(team)
        )
        candidates.append((score, team))
    candidates.sort(key=lambda item: item[0], reverse=True)
    if not candidates or candidates[0][0] < 0.88:
        return None
    if len(candidates) > 1 and candidates[0][0] - candidates[1][0] < 0.025:
        return None
    return candidates[0][1]


def _team_aliases(team):
    name = _normalize(team["name"])
    branch = _normalize(team["branch"])
    category = _normalize(team["category"])
    branch_aliases = {
        "femenil": ("fem", "femenil"),
        "varonil": ("var", "varonil"),
        "mixto": ("mix", "mixto"),
    }.get(branch, (branch,))
    category_aliases = (category, category[1:]) if category.startswith("u") else (category,)
    aliases = {name}
    for category_alias in category_aliases:
        aliases.add(f"{name} {category_alias}")
        for branch_alias in branch_aliases:
            aliases.add(f"{name} {branch_alias} {category_alias}")
    return aliases


def _week(value):
    match = re.search(r"(?:semana|jornada)\s*(\d+)", _normalize(value))
    return int(match.group(1)) if match else None


def _time_value(value):
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    text = _text(value)
    for pattern in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, pattern).time()
        except ValueError:
            pass
    return None


def _text(value):
    return "" if value is None else str(value).strip()


def _normalize(value):
    text = unicodedata.normalize("NFKD", _text(value).lower())
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())
