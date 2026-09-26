"""Parse roster imports from CSV and XLSX uploads."""

from csv import DictReader, Sniffer
from io import BytesIO, StringIO
import unicodedata

from openpyxl import load_workbook
from pydantic import ValidationError

from models import PlayerCreate


class RosterImportError(Exception):
    """Raised when an uploaded roster cannot be imported safely."""


HEADER_ALIASES = {
    "name": {"name", "nombre", "jugador", "player"},
    "curp": {"curp"},
    "age": {"age", "edad"},
    "jersey_number": {
        "numero",
        "numero_jugador",
        "numero_de_jugador",
        "dorsal",
        "jersey",
        "jersey_number",
    },
}
MAX_ROSTER_ROWS = 100


def parse_roster_file(filename, content):
    """Return validated players from a CSV or XLSX roster upload."""
    lower_name = filename.lower() if filename else ""
    if lower_name.endswith(".csv"):
        rows = _read_csv(content)
    elif lower_name.endswith(".xlsx"):
        rows = _read_xlsx(content)
    else:
        raise RosterImportError("Se requiere un archivo .csv o .xlsx")

    return _validate_rows(rows)


def _read_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise RosterImportError("El CSV debe estar codificado en UTF-8") from error

    try:
        dialect = Sniffer().sniff(text[:4096], delimiters=",;")
    except Exception:
        dialect = "excel"

    reader = DictReader(StringIO(text), dialect=dialect)
    if reader.fieldnames is None:
        raise RosterImportError("El archivo no contiene encabezados")

    header_map = _map_headers(reader.fieldnames)
    return [
        {
            "row_number": index,
            "values": {
                field: row.get(reader.fieldnames[column_index], "")
                for field, column_index in header_map.items()
            },
        }
        for index, row in enumerate(reader, start=2)
    ]


def _read_xlsx(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise RosterImportError("No se pudo leer el archivo .xlsx") from error

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise RosterImportError("El archivo no contiene encabezados")

        header_map = _map_headers(headers)
        parsed_rows = []
        for index, row in enumerate(rows, start=2):
            parsed_rows.append({
                "row_number": index,
                "values": {
                    field: _cell_to_text(row[column_index])
                    for field, column_index in header_map.items()
                },
            })
        return parsed_rows
    finally:
        workbook.close()


def _map_headers(headers):
    normalized = {
        _normalize_header(header): index
        for index, header in enumerate(headers)
        if header is not None and str(header).strip()
    }

    mapped = {}
    missing = []
    for field, aliases in HEADER_ALIASES.items():
        match = next((alias for alias in aliases if alias in normalized), None)
        if match is None:
            missing.append(field)
        else:
            mapped[field] = normalized[match]

    if missing:
        raise RosterImportError(
            "Faltan columnas requeridas: nombre, curp, edad y numero"
        )

    return mapped


def _validate_rows(rows):
    players = []
    seen_curps = {}
    seen_numbers = {}
    for row in rows:
        values = {
            key: _cell_to_text(value)
            for key, value in row["values"].items()
        }
        if not any(values.values()):
            continue
        try:
            player = PlayerCreate.model_validate(values)
        except ValidationError as error:
            raise RosterImportError(
                f"Fila {row['row_number']}: revisa nombre, CURP, edad y numero"
            ) from error

        if player.curp in seen_curps:
            raise RosterImportError(
                f"Fila {row['row_number']}: la CURP ya aparece en la fila "
                f"{seen_curps[player.curp]}"
            )
        if player.jersey_number in seen_numbers:
            raise RosterImportError(
                f"Fila {row['row_number']}: el numero {player.jersey_number} "
                f"ya aparece en la fila {seen_numbers[player.jersey_number]}"
            )

        seen_curps[player.curp] = row["row_number"]
        seen_numbers[player.jersey_number] = row["row_number"]
        players.append(player)

        if len(players) > MAX_ROSTER_ROWS:
            raise RosterImportError(
                f"El archivo no puede contener mas de {MAX_ROSTER_ROWS} jugadores"
            )

    if not players:
        raise RosterImportError("El archivo no contiene jugadores")

    return players


def _normalize_header(value):
    text = _cell_to_text(value).lower().replace("#", "numero")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return "_".join(text.replace(".", " ").replace("-", " ").split())


def _cell_to_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
